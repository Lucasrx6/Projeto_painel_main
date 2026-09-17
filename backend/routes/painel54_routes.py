"""
Painel 54 - Gestao de Transportes de Material
Subsistema de Transporte de Material — HAC
"""
from flask import Blueprint, jsonify, request, send_from_directory, send_file, session, current_app
from datetime import datetime, date
from decimal import Decimal
from backend.database import get_db_cursor
from backend.middleware.decorators import login_required, panel_permission_required
from backend.cache import cache_route, cache_delete_pattern
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule

painel54_bp = Blueprint('painel54', __name__)

# Whitelists para UPDATEs dinâmicos — nunca iterar sobre request diretamente
_CAMPOS_MOTORISTA  = ('nome', 'matricula', 'turno')
_CAMPOS_TIPO_CARGA = ('nome', 'icone', 'cor', 'ordem', 'requer_lista_itens', 'requer_assinatura',
                      'requer_foto', 'foto_obrigatoria')
_CAMPOS_DESTINO    = ('nome', 'tipo_carga_id', 'ordem', 'km_distancia')
_CAMPOS_ORIGEM     = ('nome', 'ordem', 'km_distancia')
_CAMPOS_VEICULO    = ('tipo', 'placa', 'descricao', 'km_max_dia')

# Cores Excel
_X_HAC      = '9B1C24'
_X_VERDE    = '28A745'
_X_VERMELHO = 'DC3545'
_X_LARANJA  = 'E67E00'
_X_AZUL     = '17A2B8'
_X_ROXO     = '6F42C1'
_X_BRANCO   = 'FFFFFF'
_X_ZEBRA    = 'FEF0F0'
_X_STATUS   = {
    'entregue':      _X_VERDE,
    'cancelado':     _X_VERMELHO,
    'aguardando':    _X_HAC,
    'aceito':        _X_AZUL,
    'em_transporte': _X_LARANJA,
}


def _serial(row):
    out = {}
    for k, v in row.items():
        if isinstance(v, (datetime, date)):
            out[k] = v.isoformat()
        elif isinstance(v, Decimal):
            out[k] = float(v)
        else:
            out[k] = v
    return out


def _periodo_where(req):
    data_inicio = req.args.get('data_inicio', '').strip()
    data_fim    = req.args.get('data_fim', '').strip()
    if data_inicio and data_fim:
        return (
            ["criado_em >= %s::date", "criado_em < (%s::date + INTERVAL '1 day')"],
            [data_inicio, data_fim],
        )
    dias = min(int(req.args.get('dias', 30)), 365)
    return (["criado_em >= NOW() - (%s || ' days')::INTERVAL"], [str(dias)])


def _cfg_atualizar(tabela, campos_wl, dados, rec_id, mensagem, log_tag):
    try:
        with get_db_cursor(use_dict_cursor=False) as cursor:
            fields, params = [], []
            for campo in campos_wl:
                if campo not in dados:
                    continue
                val = dados[campo]
                if isinstance(val, str):
                    val = val.strip() or None
                if campo == 'nome' and not val:
                    return jsonify({'success': False, 'error': 'Nome nao pode ser vazio'}), 400
                fields.append(campo + ' = %s')
                params.append(val)
            if 'ativo' in dados:
                fields.append('ativo = %s')
                params.append(bool(dados['ativo']))
            if not fields:
                return jsonify({'success': False, 'error': 'Nada para atualizar'}), 400
            fields.append('atualizado_em = NOW()')
            params.append(rec_id)
            cursor.execute(
                'UPDATE ' + tabela + ' SET ' + ', '.join(fields) + ' WHERE id = %s',
                params
            )
        return jsonify({'success': True, 'message': mensagem})
    except Exception as e:
        current_app.logger.error('Erro %s: %s', log_tag, e, exc_info=True)
        return jsonify({'success': False, 'error': 'Erro ao atualizar'}), 500


# helpers Excel
def _x_hdr(cell, cor=_X_HAC):
    cell.font      = Font(bold=True, color=_X_BRANCO, size=11)
    cell.fill      = PatternFill('solid', fgColor=cor)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    _x_borda(cell)


def _x_borda(cell):
    s = Side(style='thin', color='CCCCCC')
    cell.border = Border(left=s, right=s, top=s, bottom=s)


def _x_autowidth(ws, min_w=10, max_w=45):
    for col in ws.columns:
        letra = get_column_letter(col[0].column)
        w = max((len(str(c.value or '')) for c in col), default=0)
        ws.column_dimensions[letra].width = min(max(w + 3, min_w), max_w)


def _x_titulo(ws, texto, ncols, row=1):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=texto)
    c.font      = Font(bold=True, color=_X_BRANCO, size=13)
    c.fill      = PatternFill('solid', fgColor=_X_HAC)
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[row].height = 28


def _x_cor_tempo(ws, col_letra, row_ini, row_fim):
    ws.conditional_formatting.add(
        col_letra + str(row_ini) + ':' + col_letra + str(row_fim),
        ColorScaleRule(
            start_type='num', start_value=0,  start_color='63BE7B',
            mid_type='num',   mid_value=20,   mid_color='FFEB84',
            end_type='num',   end_value=60,   end_color='F8696B',
        ),
    )


# =========================================================
# PÁGINA
# =========================================================

@painel54_bp.route('/painel/painel54')
@login_required
@panel_permission_required('painel54')
def painel54():
    return send_from_directory('paineis/painel54', 'index.html')


# =========================================================
# DASHBOARD
# =========================================================

@painel54_bp.route('/api/paineis/painel54/dashboard')
@login_required
@panel_permission_required('painel54')
@cache_route(ttl=30, key_prefix='painel54:dashboard', vary_by_user=False)
def api_painel54_dashboard():
    try:
        with get_db_cursor() as cursor:
            cursor.execute("""
                SELECT
                    COUNT(*) FILTER (WHERE status = 'aguardando')                              AS aguardando,
                    COUNT(*) FILTER (WHERE status = 'aceito')                                  AS aceito,
                    COUNT(*) FILTER (WHERE status = 'em_transporte')                           AS em_transporte,
                    COUNT(*) FILTER (WHERE status = 'entregue'  AND criado_em >= CURRENT_DATE) AS entregues_hoje,
                    COUNT(*) FILTER (WHERE status = 'cancelado' AND criado_em >= CURRENT_DATE) AS cancelados_hoje,
                    COUNT(*) FILTER (WHERE criado_em >= CURRENT_DATE)                          AS total_hoje,
                    COUNT(*) FILTER (WHERE prioridade = 'urgente' AND status = 'aguardando')   AS urgentes_aguardando,
                    ROUND(AVG(EXTRACT(EPOCH FROM (dt_entrega - criado_em)) / 60)
                        FILTER (WHERE status = 'entregue' AND criado_em >= CURRENT_DATE
                                AND dt_entrega IS NOT NULL
                                AND EXTRACT(EPOCH FROM (dt_entrega - criado_em)) / 60 <= 300), 1) AS tempo_medio_total_hoje,
                    ROUND(AVG(EXTRACT(EPOCH FROM (dt_aceite - criado_em)) / 60)
                        FILTER (WHERE dt_aceite IS NOT NULL AND criado_em >= CURRENT_DATE
                                AND EXTRACT(EPOCH FROM (dt_aceite - criado_em)) / 60 <= 300), 1) AS tempo_medio_aceite_hoje,
                    ROUND(AVG(EXTRACT(EPOCH FROM (dt_entrega - dt_inicio_transporte)) / 60)
                        FILTER (WHERE status = 'entregue' AND dt_inicio_transporte IS NOT NULL
                                AND dt_entrega IS NOT NULL AND criado_em >= CURRENT_DATE
                                AND EXTRACT(EPOCH FROM (dt_entrega - criado_em)) / 60 <= 300), 1) AS tempo_medio_entrega_hoje
                FROM transporte_material_solicitacoes
                WHERE status IN ('aguardando', 'aceito', 'em_transporte')
                   OR criado_em >= CURRENT_DATE
            """)
            stats = dict(cursor.fetchone() or {})
            for k, v in stats.items():
                if v is not None and hasattr(v, '__float__'):
                    stats[k] = float(v)

            cursor.execute("""
                SELECT
                    s.id, s.nr_protocolo, s.tipo_carga_nome,
                    tc.icone AS tipo_carga_icone, tc.cor AS tipo_carga_cor,
                    s.descricao, s.setor_origem_nome, s.destino_nome,
                    s.prioridade, s.status, s.solicitante_nome, s.motorista_nome,
                    s.criado_em, s.dt_aceite, s.dt_inicio_transporte,
                    ROUND(EXTRACT(EPOCH FROM (NOW() - s.criado_em)) / 60, 1) AS minutos_espera
                FROM transporte_material_solicitacoes s
                JOIN transporte_material_tipos_carga tc ON tc.id = s.tipo_carga_id
                WHERE s.status IN ('aguardando', 'aceito', 'em_transporte')
                ORDER BY
                    CASE s.prioridade WHEN 'urgente' THEN 0 ELSE 1 END,
                    s.criado_em ASC
            """)
            ativos = []
            for row in cursor.fetchall():
                c = _serial(dict(row))
                if c.get('minutos_espera') is not None:
                    c['minutos_espera'] = float(c['minutos_espera'])
                ativos.append(c)

        return jsonify({'success': True, 'stats': stats, 'ativos': ativos,
                        'timestamp': datetime.now().isoformat()})
    except Exception as e:
        current_app.logger.error('Erro dashboard painel54: %s', e, exc_info=True)
        return jsonify({'success': False, 'error': 'Erro ao buscar dados'}), 500


# =========================================================
# LISTAGEM COM FILTROS
# =========================================================

@painel54_bp.route('/api/paineis/painel54/chamados')
@login_required
@panel_permission_required('painel54')
@cache_route(ttl=30, key_prefix='painel54:chamados', vary_by_user=False, vary_by_query=True)
def api_painel54_chamados():
    setor        = request.args.get('setor', '').strip()
    motorista_id = request.args.get('motorista_id', '').strip()
    tipo_id      = request.args.get('tipo_id', '').strip()
    status       = request.args.get('status', '').strip()
    prioridade   = request.args.get('prioridade', '').strip()

    try:
        with get_db_cursor() as cursor:
            where, params = _periodo_where(request)
            if setor:        where.append('setor_origem_nome ILIKE %s'); params.append('%' + setor + '%')
            if motorista_id: where.append('motorista_id = %s');          params.append(motorista_id)
            if tipo_id:      where.append('tipo_carga_id = %s');         params.append(tipo_id)
            if status:       where.append('status = %s');                params.append(status)
            if prioridade:   where.append('prioridade = %s');            params.append(prioridade)

            cursor.execute("""
                SELECT
                    id, nr_protocolo, tipo_carga_nome,
                    descricao, setor_origem_nome, destino_nome, destino_complemento,
                    prioridade, status, solicitante_nome, motorista_nome, observacao,
                    nm_destinatario, observacao_entrega,
                    CASE WHEN assinatura_id IS NOT NULL THEN TRUE ELSE FALSE END AS tem_assinatura,
                    criado_em, dt_aceite, dt_inicio_transporte, dt_entrega,
                    dt_cancelamento, motivo_cancelamento,
                    CASE WHEN dt_aceite IS NOT NULL
                         THEN ROUND(EXTRACT(EPOCH FROM (dt_aceite - criado_em)) / 60, 1)
                    END AS t_aceite_min,
                    CASE WHEN dt_entrega IS NOT NULL AND dt_inicio_transporte IS NOT NULL
                         THEN ROUND(EXTRACT(EPOCH FROM (dt_entrega - dt_inicio_transporte)) / 60, 1)
                    END AS t_entrega_min,
                    CASE WHEN dt_entrega IS NOT NULL
                         THEN ROUND(EXTRACT(EPOCH FROM (dt_entrega - criado_em)) / 60, 1)
                    END AS t_total_min
                FROM transporte_material_solicitacoes
                WHERE """ + ' AND '.join(where) + """
                ORDER BY criado_em DESC
                LIMIT 500
            """, params)

            chamados = [_serial(dict(r)) for r in cursor.fetchall()]
        return jsonify({'success': True, 'chamados': chamados, 'total': len(chamados)})
    except Exception as e:
        current_app.logger.error('Erro chamados painel54: %s', e, exc_info=True)
        return jsonify({'success': False, 'error': 'Erro ao buscar chamados'}), 500


# =========================================================
# CANCELAR (admin)
# =========================================================

@painel54_bp.route('/api/paineis/painel54/chamados/<int:chamado_id>/cancelar', methods=['PUT'])
@login_required
@panel_permission_required('painel54')
def api_painel54_cancelar(chamado_id):
    motivo = ((request.get_json() or {}).get('motivo') or '').strip()
    if len(motivo) < 10:
        return jsonify({'success': False,
                        'error': 'O motivo deve ter pelo menos 10 caracteres'}), 400
    try:
        with get_db_cursor() as cursor:
            cursor.execute(
                "SELECT status FROM transporte_material_solicitacoes WHERE id = %s",
                (chamado_id,)
            )
            chamado = cursor.fetchone()
            if not chamado:
                return jsonify({'success': False, 'error': 'Chamado nao encontrado'}), 404
            if chamado['status'] in ('entregue', 'cancelado'):
                return jsonify({'success': False,
                                'error': 'Chamado nao pode ser cancelado no status: ' + chamado['status']}), 400
            cursor.execute("""
                UPDATE transporte_material_solicitacoes
                SET status = 'cancelado', dt_cancelamento = NOW(),
                    motivo_cancelamento = %s, atualizado_em = NOW()
                WHERE id = %s
            """, ('[Cancelado pela Gestao] ' + motivo, chamado_id))

        cache_delete_pattern('painel54:*')
        return jsonify({'success': True, 'message': 'Chamado cancelado administrativamente'})
    except Exception as e:
        current_app.logger.error('Erro cancelar painel54: %s', e, exc_info=True)
        return jsonify({'success': False, 'error': 'Erro ao cancelar chamado'}), 500


# =========================================================
# VER ASSINATURA
# =========================================================

@painel54_bp.route('/api/paineis/painel54/chamados/<int:chamado_id>/assinatura')
@login_required
@panel_permission_required('painel54')
def api_painel54_assinatura(chamado_id):
    try:
        with get_db_cursor() as cursor:
            cursor.execute("""
                SELECT s.nr_protocolo, s.nm_destinatario, s.dt_entrega,
                       a.assinatura_img, a.nm_signatario, a.criado_em AS dt_assinatura,
                       a.coletado_por_nome, a.coletado_por_matricula
                FROM transporte_material_solicitacoes s
                LEFT JOIN assinaturas_digitais a ON a.id = s.assinatura_id
                WHERE s.id = %s
            """, (chamado_id,))
            row = cursor.fetchone()
        if not row:
            return jsonify({'success': False, 'error': 'Chamado nao encontrado'}), 404
        return jsonify({'success': True, 'assinatura': _serial(dict(row))})
    except Exception as e:
        current_app.logger.error('Erro assinatura painel54: %s', e, exc_info=True)
        return jsonify({'success': False, 'error': 'Erro ao buscar assinatura'}), 500


# =========================================================
# ANALYTICS POR SETOR
# =========================================================

@painel54_bp.route('/api/paineis/painel54/por-setor')
@login_required
@panel_permission_required('painel54')
@cache_route(ttl=120, key_prefix='painel54:por-setor', vary_by_user=False, vary_by_query=True)
def api_painel54_por_setor():
    try:
        with get_db_cursor() as cursor:
            where, params = _periodo_where(request)
            cursor.execute("""
                SELECT
                    setor_origem_nome                                                          AS setor,
                    COUNT(*)                                                                   AS total,
                    COUNT(*) FILTER (WHERE status = 'entregue')                               AS entregues,
                    COUNT(*) FILTER (WHERE status = 'cancelado')                              AS cancelados,
                    COUNT(*) FILTER (WHERE prioridade = 'urgente')                            AS urgentes,
                    ROUND(AVG(EXTRACT(EPOCH FROM (dt_aceite - criado_em)) / 60)
                        FILTER (WHERE dt_aceite IS NOT NULL), 1)                              AS t_aceite_min,
                    ROUND(AVG(EXTRACT(EPOCH FROM (dt_entrega - criado_em)) / 60)
                        FILTER (WHERE status = 'entregue' AND dt_entrega IS NOT NULL), 1)     AS t_total_min
                FROM transporte_material_solicitacoes
                WHERE """ + ' AND '.join(where) + """
                GROUP BY setor_origem_nome
                ORDER BY total DESC
            """, params)
            setores = [_serial(dict(r)) for r in cursor.fetchall()]
        return jsonify({'success': True, 'setores': setores})
    except Exception as e:
        current_app.logger.error('Erro por-setor painel54: %s', e, exc_info=True)
        return jsonify({'success': False, 'error': 'Erro ao buscar dados'}), 500


# =========================================================
# ANALYTICS POR MOTORISTA
# =========================================================

@painel54_bp.route('/api/paineis/painel54/por-motorista')
@login_required
@panel_permission_required('painel54')
@cache_route(ttl=120, key_prefix='painel54:por-motorista', vary_by_user=False, vary_by_query=True)
def api_painel54_por_motorista():
    try:
        with get_db_cursor() as cursor:
            where, params = _periodo_where(request)
            where.append('motorista_nome IS NOT NULL')
            cursor.execute("""
                SELECT
                    motorista_nome                                                             AS motorista,
                    COUNT(*)                                                                   AS total,
                    COUNT(*) FILTER (WHERE status = 'entregue')                               AS entregues,
                    COUNT(*) FILTER (WHERE status = 'cancelado')                              AS cancelados,
                    COUNT(*) FILTER (WHERE prioridade = 'urgente')                            AS urgentes,
                    ROUND(AVG(EXTRACT(EPOCH FROM (dt_aceite - criado_em)) / 60)
                        FILTER (WHERE dt_aceite IS NOT NULL), 1)                              AS t_aceite_min,
                    ROUND(AVG(EXTRACT(EPOCH FROM (dt_entrega - dt_inicio_transporte)) / 60)
                        FILTER (WHERE status = 'entregue' AND dt_inicio_transporte IS NOT NULL
                                AND dt_entrega IS NOT NULL), 1)                               AS t_entrega_min,
                    ROUND(AVG(EXTRACT(EPOCH FROM (dt_entrega - criado_em)) / 60)
                        FILTER (WHERE status = 'entregue' AND dt_entrega IS NOT NULL), 1)     AS t_total_min
                FROM transporte_material_solicitacoes
                WHERE """ + ' AND '.join(where) + """
                GROUP BY motorista_nome
                ORDER BY entregues DESC
            """, params)
            motoristas = [_serial(dict(r)) for r in cursor.fetchall()]
        return jsonify({'success': True, 'motoristas': motoristas})
    except Exception as e:
        current_app.logger.error('Erro por-motorista painel54: %s', e, exc_info=True)
        return jsonify({'success': False, 'error': 'Erro ao buscar dados'}), 500


# =========================================================
# EXPORTAR EXCEL
# =========================================================

@painel54_bp.route('/api/paineis/painel54/exportar')
@login_required
@panel_permission_required('painel54')
def api_painel54_exportar():
    status_f    = request.args.get('status', '').strip()
    prioridade  = request.args.get('prioridade', '').strip()
    setor       = request.args.get('setor', '').strip()
    data_inicio = request.args.get('data_inicio', '').strip()
    data_fim    = request.args.get('data_fim', '').strip()

    try:
        with get_db_cursor() as cursor:
            where, params = _periodo_where(request)
            if status_f:   where.append('status = %s');                params.append(status_f)
            if prioridade: where.append('prioridade = %s');            params.append(prioridade)
            if setor:      where.append('setor_origem_nome ILIKE %s'); params.append('%' + setor + '%')
            where_sql = ' AND '.join(where)

            cursor.execute("""
                SELECT
                    id, nr_protocolo, tipo_carga_nome,
                    descricao, setor_origem_nome, destino_nome, destino_complemento,
                    prioridade, status, solicitante_nome, motorista_nome,
                    nm_destinatario, observacao,
                    TO_CHAR(criado_em,            'DD/MM/YYYY HH24:MI') AS criado_em,
                    TO_CHAR(dt_aceite,            'DD/MM/YYYY HH24:MI') AS dt_aceite,
                    TO_CHAR(dt_inicio_transporte, 'DD/MM/YYYY HH24:MI') AS dt_inicio_transporte,
                    TO_CHAR(dt_entrega,           'DD/MM/YYYY HH24:MI') AS dt_entrega,
                    TO_CHAR(dt_cancelamento,      'DD/MM/YYYY HH24:MI') AS dt_cancelamento,
                    motivo_cancelamento,
                    CASE WHEN assinatura_id IS NOT NULL THEN 'Sim' ELSE 'Nao' END AS assinado,
                    CASE WHEN dt_aceite IS NOT NULL
                         THEN ROUND(EXTRACT(EPOCH FROM (dt_aceite - criado_em)) / 60, 1) END AS t_aceite_min,
                    CASE WHEN dt_entrega IS NOT NULL AND dt_inicio_transporte IS NOT NULL
                         THEN ROUND(EXTRACT(EPOCH FROM (dt_entrega - dt_inicio_transporte)) / 60, 1) END AS t_entrega_min,
                    CASE WHEN dt_entrega IS NOT NULL
                         THEN ROUND(EXTRACT(EPOCH FROM (dt_entrega - criado_em)) / 60, 1)
                         WHEN status = 'cancelado' AND dt_cancelamento IS NOT NULL
                         THEN ROUND(EXTRACT(EPOCH FROM (dt_cancelamento - criado_em)) / 60, 1)
                    END AS t_total_min
                FROM transporte_material_solicitacoes WHERE """ + where_sql + """
                ORDER BY criado_em DESC
            """, params)
            chamados = [dict(r) for r in cursor.fetchall()]

            cursor.execute("""
                SELECT
                    COALESCE(motorista_nome,'(nao atribuido)') AS motorista,
                    COUNT(*) AS total,
                    COUNT(*) FILTER (WHERE status='entregue')   AS entregues,
                    COUNT(*) FILTER (WHERE status='cancelado')  AS cancelados,
                    COUNT(*) FILTER (WHERE prioridade='urgente') AS urgentes,
                    ROUND(AVG(CASE WHEN dt_aceite IS NOT NULL
                        THEN EXTRACT(EPOCH FROM (dt_aceite - criado_em))/60 END)::numeric,1) AS media_aceite,
                    ROUND(AVG(CASE WHEN dt_entrega IS NOT NULL
                        THEN EXTRACT(EPOCH FROM (dt_entrega - criado_em))/60 END)::numeric,1) AS media_total
                FROM transporte_material_solicitacoes WHERE """ + where_sql + """
                GROUP BY motorista_nome ORDER BY total DESC
            """, params)
            por_motorista = [dict(r) for r in cursor.fetchall()]

            cursor.execute("""
                SELECT
                    COALESCE(setor_origem_nome,'(sem setor)') AS setor,
                    COUNT(*) AS total,
                    COUNT(*) FILTER (WHERE status='entregue')   AS entregues,
                    COUNT(*) FILTER (WHERE status='cancelado')  AS cancelados,
                    COUNT(*) FILTER (WHERE prioridade='urgente') AS urgentes
                FROM transporte_material_solicitacoes WHERE """ + where_sql + """
                GROUP BY setor_origem_nome ORDER BY total DESC LIMIT 30
            """, params)
            por_setor = [dict(r) for r in cursor.fetchall()]

        now = datetime.now()
        if data_inicio and data_fim:
            di = datetime.strptime(data_inicio, '%Y-%m-%d').strftime('%d/%m/%Y')
            df = datetime.strptime(data_fim,    '%Y-%m-%d').strftime('%d/%m/%Y')
            filtros_txt = 'Periodo: ' + di + ' a ' + df
        else:
            filtros_txt = 'Ultimos ' + str(int(request.args.get('dias', 30))) + ' dia(s)'
        if status_f:   filtros_txt += '  |  Status: ' + status_f
        if prioridade: filtros_txt += '  |  Prioridade: ' + prioridade
        if setor:      filtros_txt += '  |  Setor: ' + setor

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        # ── Aba 1: Chamados ───────────────────────────────────
        ws1 = wb.create_sheet('Chamados')
        ws1.sheet_view.showGridLines = False
        ws1.freeze_panes = 'A3'
        _x_titulo(ws1,
                  'TRANSPORTES DE MATERIAL — HAC — ' + now.strftime('%d/%m/%Y %H:%M') + '  |  ' + filtros_txt,
                  24)
        hdrs1 = [
            '#', 'Protocolo', 'Tipo Carga', 'Descricao', 'Origem', 'Destino', 'Compl. Destino',
            'Prioridade', 'Status', 'Solicitante', 'Motorista', 'Destinatario', 'Observacao',
            'Criado Em', 'Aceito Em', 'Ini. Transporte', 'Entregue Em', 'Cancelado Em',
            'Motivo Cancelamento', 'Assinado?', 'T.Aceite(min)', 'T.Entrega(min)', 'T.Total(min)',
        ]
        for j, h in enumerate(hdrs1, 1):
            _x_hdr(ws1.cell(row=2, column=j, value=h))
        ws1.row_dimensions[2].height = 22

        keys1 = [
            'id', 'nr_protocolo', 'tipo_carga_nome', 'descricao', 'setor_origem_nome',
            'destino_nome', 'destino_complemento', 'prioridade', 'status',
            'solicitante_nome', 'motorista_nome', 'nm_destinatario', 'observacao',
            'criado_em', 'dt_aceite', 'dt_inicio_transporte', 'dt_entrega', 'dt_cancelamento',
            'motivo_cancelamento', 'assinado', 't_aceite_min', 't_entrega_min', 't_total_min',
        ]
        for i, row in enumerate(chamados, 3):
            zebra = i % 2 == 0
            for j, key in enumerate(keys1, 1):
                v    = row.get(key)
                cell = ws1.cell(row=i, column=j, value=v)
                _x_borda(cell)
                cell.alignment = Alignment(vertical='center', wrap_text=(j in (4, 13, 19)))
                if zebra and key not in ('prioridade', 'status'):
                    cell.fill = PatternFill('solid', fgColor=_X_ZEBRA)
                if key == 'status' and v in _X_STATUS:
                    cell.font = Font(bold=True, color=_X_BRANCO)
                    cell.fill = PatternFill('solid', fgColor=_X_STATUS[v])
                elif key == 'prioridade' and v == 'urgente':
                    cell.font = Font(bold=True, color=_X_BRANCO)
                    cell.fill = PatternFill('solid', fgColor=_X_LARANJA)

        last1 = 2 + len(chamados)
        if chamados:
            for col_n in range(21, 24):
                _x_cor_tempo(ws1, get_column_letter(col_n), 3, last1)
        _x_autowidth(ws1)

        # ── Aba 2: Por Motorista ──────────────────────────────
        ws2 = wb.create_sheet('Por Motorista')
        ws2.sheet_view.showGridLines = False
        hdrs2 = ['Motorista', 'Total', 'Entregues', 'Cancelados', 'Urgentes',
                 'T.Aceite(min)', 'T.Total(min)']
        _x_titulo(ws2, 'POR MOTORISTA — ' + filtros_txt, len(hdrs2))
        for j, h in enumerate(hdrs2, 1):
            _x_hdr(ws2.cell(row=2, column=j, value=h))

        mot_cols = [
            ('motorista', None), ('total', None), ('entregues', _X_VERDE),
            ('cancelados', _X_VERMELHO), ('urgentes', _X_LARANJA),
            ('media_aceite', None), ('media_total', None),
        ]
        for i, p in enumerate(por_motorista, 3):
            zebra = i % 2 == 0
            for j, (key, cor_txt) in enumerate(mot_cols, 1):
                v    = p.get(key)
                cell = ws2.cell(row=i, column=j, value=float(v) if v is not None and j > 5 else v)
                _x_borda(cell)
                if zebra:     cell.fill = PatternFill('solid', fgColor=_X_ZEBRA)
                if j == 1:    cell.font = Font(bold=True)
                elif cor_txt: cell.font = Font(bold=True, color=cor_txt)

        last2 = 2 + len(por_motorista)
        if por_motorista:
            for col_l in ['F', 'G']:
                _x_cor_tempo(ws2, col_l, 3, last2)
            c = BarChart()
            c.type = 'bar'; c.grouping = 'clustered'
            c.title = 'Transportes por Motorista'
            c.style = 10; c.height = 12; c.width = 22
            c.add_data(Reference(ws2, min_col=2, max_col=5, min_row=2, max_row=last2), titles_from_data=True)
            c.set_categories(Reference(ws2, min_col=1, min_row=3, max_row=last2))
            for idx, cor in enumerate([_X_HAC, _X_VERDE, _X_VERMELHO, _X_LARANJA]):
                if idx < len(c.series):
                    c.series[idx].graphicalProperties.solidFill = cor
            ws2.add_chart(c, 'A' + str(last2 + 3))
        _x_autowidth(ws2)

        # ── Aba 3: Por Setor ──────────────────────────────────
        ws3 = wb.create_sheet('Por Setor')
        ws3.sheet_view.showGridLines = False
        hdrs3 = ['Setor', 'Total', 'Entregues', 'Cancelados', 'Urgentes']
        _x_titulo(ws3, 'POR SETOR — ' + filtros_txt, len(hdrs3))
        for j, h in enumerate(hdrs3, 1):
            _x_hdr(ws3.cell(row=2, column=j, value=h))

        set_cols = [
            ('setor', None), ('total', None), ('entregues', _X_VERDE),
            ('cancelados', _X_VERMELHO), ('urgentes', _X_LARANJA),
        ]
        for i, s in enumerate(por_setor, 3):
            zebra = i % 2 == 0
            for j, (key, cor_txt) in enumerate(set_cols, 1):
                cell = ws3.cell(row=i, column=j, value=s.get(key))
                _x_borda(cell)
                if zebra:     cell.fill = PatternFill('solid', fgColor=_X_ZEBRA)
                if j == 1:    cell.font = Font(bold=True)
                elif cor_txt: cell.font = Font(bold=True, color=cor_txt)

        last3 = 2 + len(por_setor)
        if por_setor:
            c2 = BarChart()
            c2.type = 'bar'; c2.grouping = 'clustered'
            c2.title = 'Chamados por Setor'
            c2.style = 10; c2.height = max(12, len(por_setor) * 0.9); c2.width = 22
            c2.add_data(Reference(ws3, min_col=2, max_col=5, min_row=2, max_row=last3), titles_from_data=True)
            c2.set_categories(Reference(ws3, min_col=1, min_row=3, max_row=last3))
            for idx, cor in enumerate([_X_HAC, _X_VERDE, _X_VERMELHO, _X_LARANJA]):
                if idx < len(c2.series):
                    c2.series[idx].graphicalProperties.solidFill = cor
            ws3.add_chart(c2, 'A' + str(last3 + 3))
        _x_autowidth(ws3)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(
            buf,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name='transportes_material_' + date.today().strftime('%Y%m%d') + '.xlsx',
        )
    except Exception as e:
        current_app.logger.error('Erro exportar painel54: %s', e, exc_info=True)
        return jsonify({'success': False, 'error': 'Erro ao exportar'}), 500


# =========================================================
# CONFIG: MOTORISTAS
# =========================================================

@painel54_bp.route('/api/paineis/painel54/config/motoristas')
@login_required
@panel_permission_required('painel54')
def api_painel54_cfg_mot_listar():
    try:
        with get_db_cursor() as cursor:
            cursor.execute("""
                SELECT id, nome, matricula, turno, ativo,
                       TO_CHAR(criado_em, 'DD/MM/YYYY') AS criado_em
                FROM transporte_material_motoristas ORDER BY nome
            """)
            return jsonify({'success': True, 'motoristas': [dict(r) for r in cursor.fetchall()]})
    except Exception as e:
        current_app.logger.error('Erro listar motoristas painel54: %s', e, exc_info=True)
        return jsonify({'success': False, 'error': 'Erro ao buscar motoristas'}), 500


@painel54_bp.route('/api/paineis/painel54/config/motoristas', methods=['POST'])
@login_required
@panel_permission_required('painel54')
def api_painel54_cfg_mot_criar():
    dados = request.get_json() or {}
    nome  = (dados.get('nome') or '').strip()
    if not nome:
        return jsonify({'success': False, 'error': 'Nome e obrigatorio'}), 400
    try:
        with get_db_cursor() as cursor:
            cursor.execute("""
                INSERT INTO transporte_material_motoristas (nome, matricula, turno, ativo, criado_em)
                VALUES (%s, %s, %s, TRUE, NOW()) RETURNING id
            """, (nome, (dados.get('matricula') or '').strip() or None, dados.get('turno', 'todos')))
            return jsonify({'success': True, 'id': cursor.fetchone()['id'],
                            'message': 'Motorista cadastrado'}), 201
    except Exception as e:
        current_app.logger.error('Erro criar motorista painel54: %s', e, exc_info=True)
        return jsonify({'success': False, 'error': 'Erro ao cadastrar motorista'}), 500


@painel54_bp.route('/api/paineis/painel54/config/motoristas/<int:motorista_id>', methods=['PUT'])
@login_required
@panel_permission_required('painel54')
def api_painel54_cfg_mot_atualizar(motorista_id):
    return _cfg_atualizar(
        'transporte_material_motoristas', _CAMPOS_MOTORISTA,
        request.get_json() or {}, motorista_id,
        'Motorista atualizado', 'atualizar motorista painel54'
    )


# =========================================================
# CONFIG: TIPOS DE CARGA
# =========================================================

@painel54_bp.route('/api/paineis/painel54/config/tipos-carga')
@login_required
@panel_permission_required('painel54')
def api_painel54_cfg_tipo_listar():
    try:
        with get_db_cursor() as cursor:
            cursor.execute("""
                SELECT id, nome, icone, cor, requer_lista_itens, requer_assinatura,
                       requer_foto, foto_obrigatoria, ativo, ordem
                FROM transporte_material_tipos_carga ORDER BY ordem, nome
            """)
            return jsonify({'success': True, 'tipos': [dict(r) for r in cursor.fetchall()]})
    except Exception as e:
        current_app.logger.error('Erro listar tipos painel54: %s', e, exc_info=True)
        return jsonify({'success': False, 'error': 'Erro ao buscar tipos'}), 500


@painel54_bp.route('/api/paineis/painel54/config/tipos-carga', methods=['POST'])
@login_required
@panel_permission_required('painel54')
def api_painel54_cfg_tipo_criar():
    dados = request.get_json() or {}
    nome  = (dados.get('nome') or '').strip()
    if not nome:
        return jsonify({'success': False, 'error': 'Nome e obrigatorio'}), 400
    try:
        with get_db_cursor() as cursor:
            cursor.execute("""
                INSERT INTO transporte_material_tipos_carga
                    (nome, icone, cor, requer_lista_itens, requer_assinatura,
                     requer_foto, foto_obrigatoria, ativo, ordem)
                VALUES (%s, %s, %s, %s, %s, %s, %s, TRUE, %s) RETURNING id
            """, (
                nome,
                dados.get('icone', 'fa-box'),
                dados.get('cor', '#6c757d'),
                bool(dados.get('requer_lista_itens', False)),
                bool(dados.get('requer_assinatura', False)),
                bool(dados.get('requer_foto', False)),
                bool(dados.get('foto_obrigatoria', False)),
                dados.get('ordem', 0)
            ))
            return jsonify({'success': True, 'id': cursor.fetchone()['id']}), 201
    except Exception as e:
        current_app.logger.error('Erro criar tipo painel54: %s', e, exc_info=True)
        return jsonify({'success': False, 'error': 'Erro ao criar tipo'}), 500


@painel54_bp.route('/api/paineis/painel54/config/tipos-carga/<int:tipo_id>', methods=['PUT'])
@login_required
@panel_permission_required('painel54')
def api_painel54_cfg_tipo_atualizar(tipo_id):
    return _cfg_atualizar(
        'transporte_material_tipos_carga', _CAMPOS_TIPO_CARGA,
        request.get_json() or {}, tipo_id,
        'Tipo atualizado', 'atualizar tipo painel54'
    )


# =========================================================
# CONFIG: DESTINOS
# =========================================================

@painel54_bp.route('/api/paineis/painel54/config/destinos')
@login_required
@panel_permission_required('painel54')
def api_painel54_cfg_dest_listar():
    tipo_id = request.args.get('tipo_id', '').strip()
    try:
        with get_db_cursor() as cursor:
            if tipo_id:
                cursor.execute("""
                    SELECT d.id, d.nome, d.tipo_carga_id, t.nome AS tipo_nome,
                           d.ativo, d.ordem, d.km_distancia
                    FROM transporte_material_destinos d
                    LEFT JOIN transporte_material_tipos_carga t ON t.id = d.tipo_carga_id
                    WHERE d.tipo_carga_id = %s ORDER BY d.ordem, d.nome
                """, (tipo_id,))
            else:
                cursor.execute("""
                    SELECT d.id, d.nome, d.tipo_carga_id, t.nome AS tipo_nome,
                           d.ativo, d.ordem, d.km_distancia
                    FROM transporte_material_destinos d
                    LEFT JOIN transporte_material_tipos_carga t ON t.id = d.tipo_carga_id
                    ORDER BY COALESCE(t.ordem, 999), d.ordem, d.nome
                """)
            return jsonify({'success': True, 'destinos': [dict(r) for r in cursor.fetchall()]})
    except Exception as e:
        current_app.logger.error('Erro listar destinos painel54: %s', e, exc_info=True)
        return jsonify({'success': False, 'error': 'Erro ao buscar destinos'}), 500


@painel54_bp.route('/api/paineis/painel54/config/destinos', methods=['POST'])
@login_required
@panel_permission_required('painel54')
def api_painel54_cfg_dest_criar():
    dados = request.get_json() or {}
    nome  = (dados.get('nome') or '').strip()
    if not nome:
        return jsonify({'success': False, 'error': 'Nome e obrigatorio'}), 400
    try:
        with get_db_cursor() as cursor:
            km = dados.get('km_distancia')
            try:
                km = float(km) if km not in (None, '') else None
            except (ValueError, TypeError):
                km = None
            cursor.execute("""
                INSERT INTO transporte_material_destinos (nome, tipo_carga_id, ativo, ordem, km_distancia)
                VALUES (%s, %s, TRUE, %s, %s) RETURNING id
            """, (nome, dados.get('tipo_carga_id') or None, dados.get('ordem', 0), km))
            return jsonify({'success': True, 'id': cursor.fetchone()['id']}), 201
    except Exception as e:
        current_app.logger.error('Erro criar destino painel54: %s', e, exc_info=True)
        return jsonify({'success': False, 'error': 'Erro ao criar destino'}), 500


@painel54_bp.route('/api/paineis/painel54/config/destinos/<int:destino_id>', methods=['PUT'])
@login_required
@panel_permission_required('painel54')
def api_painel54_cfg_dest_atualizar(destino_id):
    return _cfg_atualizar(
        'transporte_material_destinos', _CAMPOS_DESTINO,
        request.get_json() or {}, destino_id,
        'Destino atualizado', 'atualizar destino painel54'
    )


# =========================================================
# CONFIG: ORIGENS
# =========================================================

@painel54_bp.route('/api/paineis/painel54/config/origens')
@login_required
@panel_permission_required('painel54')
def api_painel54_cfg_orig_listar():
    try:
        with get_db_cursor() as cursor:
            cursor.execute("""
                SELECT id, nome, ativo, ordem, km_distancia
                FROM transporte_material_origens ORDER BY ordem, nome
            """)
            return jsonify({'success': True, 'origens': [dict(r) for r in cursor.fetchall()]})
    except Exception as e:
        current_app.logger.error('Erro listar origens painel54: %s', e, exc_info=True)
        return jsonify({'success': False, 'error': 'Erro ao buscar origens'}), 500


@painel54_bp.route('/api/paineis/painel54/config/origens', methods=['POST'])
@login_required
@panel_permission_required('painel54')
def api_painel54_cfg_orig_criar():
    dados = request.get_json() or {}
    nome  = (dados.get('nome') or '').strip()
    if not nome:
        return jsonify({'success': False, 'error': 'Nome e obrigatorio'}), 400
    try:
        with get_db_cursor() as cursor:
            km = dados.get('km_distancia')
            try:
                km = float(km) if km not in (None, '') else None
            except (ValueError, TypeError):
                km = None
            cursor.execute("""
                INSERT INTO transporte_material_origens (nome, ativo, ordem, km_distancia)
                VALUES (%s, TRUE, %s, %s) RETURNING id
            """, (nome, dados.get('ordem', 0), km))
            return jsonify({'success': True, 'id': cursor.fetchone()['id']}), 201
    except Exception as e:
        current_app.logger.error('Erro criar origem painel54: %s', e, exc_info=True)
        return jsonify({'success': False, 'error': 'Erro ao criar origem'}), 500


@painel54_bp.route('/api/paineis/painel54/config/origens/<int:origem_id>', methods=['PUT'])
@login_required
@panel_permission_required('painel54')
def api_painel54_cfg_orig_atualizar(origem_id):
    return _cfg_atualizar(
        'transporte_material_origens', _CAMPOS_ORIGEM,
        request.get_json() or {}, origem_id,
        'Origem atualizada', 'atualizar origem painel54'
    )


# =========================================================
# CONFIG: VEÍCULOS
# =========================================================

@painel54_bp.route('/api/paineis/painel54/config/veiculos')
@login_required
@panel_permission_required('painel54')
def api_painel54_cfg_vei_listar():
    try:
        with get_db_cursor() as cursor:
            cursor.execute("""
                SELECT id, tipo, placa, descricao, km_max_dia, ativo,
                       TO_CHAR(criado_em, 'DD/MM/YYYY') AS criado_em
                FROM transporte_material_veiculos ORDER BY tipo, placa
            """)
            veiculos = []
            for row in cursor.fetchall():
                v = dict(row)
                if v.get('km_max_dia') is not None:
                    v['km_max_dia'] = float(v['km_max_dia'])
                veiculos.append(v)
        return jsonify({'success': True, 'veiculos': veiculos})
    except Exception as e:
        current_app.logger.error('Erro listar veiculos painel54: %s', e, exc_info=True)
        return jsonify({'success': False, 'error': 'Erro ao buscar veiculos'}), 500


@painel54_bp.route('/api/paineis/painel54/config/veiculos', methods=['POST'])
@login_required
@panel_permission_required('painel54')
def api_painel54_cfg_vei_criar():
    dados = request.get_json() or {}
    tipo  = (dados.get('tipo') or 'carro').strip()
    if tipo not in ('moto', 'carro', 'van', 'caminhao', 'outro'):
        tipo = 'carro'
    placa = (dados.get('placa') or '').strip().upper() or None
    km    = dados.get('km_max_dia')
    try:
        km = float(km) if km not in (None, '') else None
    except (ValueError, TypeError):
        km = None
    try:
        with get_db_cursor() as cursor:
            cursor.execute("""
                INSERT INTO transporte_material_veiculos
                    (tipo, placa, descricao, km_max_dia, ativo, criado_em)
                VALUES (%s, %s, %s, %s, TRUE, NOW()) RETURNING id
            """, (tipo, placa, (dados.get('descricao') or '').strip() or None, km))
            return jsonify({'success': True, 'id': cursor.fetchone()['id'],
                            'message': 'Veiculo cadastrado'}), 201
    except Exception as e:
        current_app.logger.error('Erro criar veiculo painel54: %s', e, exc_info=True)
        return jsonify({'success': False, 'error': 'Erro ao cadastrar veiculo'}), 500


@painel54_bp.route('/api/paineis/painel54/config/veiculos/<int:veiculo_id>', methods=['PUT'])
@login_required
@panel_permission_required('painel54')
def api_painel54_cfg_vei_atualizar(veiculo_id):
    return _cfg_atualizar(
        'transporte_material_veiculos', _CAMPOS_VEICULO,
        request.get_json() or {}, veiculo_id,
        'Veiculo atualizado', 'atualizar veiculo painel54'
    )
